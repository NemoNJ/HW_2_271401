"""
จำลองระบบควบคุมอุณหภูมิห้องด้วยแอร์แบบ Closed Loop
Transfer Function ของ Plant (ห้อง+แอร์): G(s) = K / (tau*s + 1)

ใช้วิธีแปลง Transfer Function กลับเป็นสมการเชิงอนุพันธ์ (ODE) เอง
แล้วแก้สมการด้วย scipy.integrate.odeint แทนการใช้ ctrl.TransferFunction/ctrl.feedback

--------------------------------------------------------------------
ที่มาของสมการ (การแปลง Transfer Function -> ODE):

  Plant:  G(s) = Y(s)/U(s) = K / (tau*s + 1)
          tau*s*Y(s) + Y(s) = K*U(s)
          แปลงกลับเป็น time domain (s -> d/dt):
              tau * dy/dt + y = K * u
              dy/dt = (K*u - y) / tau              ... (1)

  Controller (P-control, unity feedback):
              u = Kp * (setpoint - y)               ... (2)

  แทน (2) ลงใน (1) จะได้ ODE ของ Closed-Loop โดยตรง:
              dy/dt = [ Kp*K*(setpoint - y) - y ] / tau
                    = [ Kp*K*setpoint - (Kp*K + 1)*y ] / tau

  สำหรับ Open-Loop เราป้อน u = step input (=1) ตรงๆ ไม่ผ่าน feedback:
              dy/dt = (K*1 - y) / tau
--------------------------------------------------------------------

การทดลอง:
  1. เปลี่ยนค่า K จาก 1 ถึง 10 (ครบทุกค่าจำนวนเต็ม, tau คงที่) ดูผลต่อ step response
  2. เปลี่ยนค่า tau เป็น 2 เท่า และ 3 เท่า (K คงที่) ดูผลต่อความเร็วในการตอบสนอง

*** เพิ่มเติมในเวอร์ชันนี้ ***
  - K ไล่ค่าครบตั้งแต่ 1 ถึง 10 ทุกจำนวนเต็ม (K_values = 1,2,3,...,10)
  - บันทึกผลลัพธ์ตัวเลข (steady-state, settling time, rise time) ลงไฟล์ .xlsx
    พร้อมจัดฟอร์แมตตาราง (หัวตาราง, ฟอนต์, เส้นขอบ, สี)
"""

import numpy as np
import matplotlib.pyplot as plt
from scipy.integrate import odeint
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

import os

OUT_DIR = '/home/jirayu/HW_271401_Control'
os.makedirs(OUT_DIR, exist_ok=True)

# -----------------------------------------------------------------
# ค่าคงที่พื้นฐาน
# -----------------------------------------------------------------
Kp = 1.0                       # Proportional gain ของตัวควบคุม (คงที่ตลอด เพื่อดู effect ของ K, tau ล้วนๆ)
tau_base = 1.0                 # ค่าเวลาคงที่พื้นฐานของห้อง (นาที)
setpoint = 1.0                 # step input ขนาด 1 หน่วย (unit step)
t = np.linspace(0, 10, 2000)   # เวลาจำลอง 0-10 นาที (จุดเยอะขึ้นเพื่อความละเอียดตอนคำนวณ rise/settling time)


# ===================================================================
# สมการเชิงอนุพันธ์ (ODE) ของระบบ
# ===================================================================
def closed_loop_ode(y, t, K, tau, Kp, setpoint):
    """dy/dt ของระบบ closed-loop: Plant + P-controller + unity feedback"""
    u = Kp * (setpoint - y)          # (2) controller คำนวณ error แล้วคูณ Kp
    dydt = (K * u - y) / tau         # (1) สมการของ plant
    return dydt


def open_loop_ode(y, t, K, tau, u):
    """dy/dt ของ plant อย่างเดียว ไม่มี feedback ป้อน u คงที่ (step) เข้าไปตรงๆ"""
    dydt = (K * u - y) / tau
    return dydt


def simulate_closed_loop(K, tau, Kp=1.0, sp=1.0, t=t):
    y0 = 0.0
    y = odeint(closed_loop_ode, y0, t, args=(K, tau, Kp, sp))
    return y.flatten()


def simulate_open_loop(K, tau, u=1.0, t=t):
    y0 = 0.0
    y = odeint(open_loop_ode, y0, t, args=(K, tau, u))
    return y.flatten()


# ===================================================================
# ฟังก์ชันช่วยคำนวณ Rise Time และ Settling Time จากข้อมูลที่จำลองได้
# (แทนที่ ctrl.step_info ซึ่งเคยใช้ library control)
# ===================================================================
def compute_rise_time(t, y, y_final):
    """เวลาที่ใช้ตอบสนองจาก 10% ถึง 90% ของค่าสุดท้าย"""
    if y_final == 0:
        return np.nan
    y10, y90 = 0.1 * y_final, 0.9 * y_final
    idx10 = np.where(y >= y10)[0]
    idx90 = np.where(y >= y90)[0]
    if len(idx10) == 0 or len(idx90) == 0:
        return np.nan
    return t[idx90[0]] - t[idx10[0]]


def compute_settling_time(t, y, y_final, tolerance=0.02):
    """เวลาที่ระบบเข้าสู่ช่วง ±2% ของค่าสุดท้ายแล้วไม่ออกมาอีก"""
    if y_final == 0:
        return np.nan
    band = tolerance * abs(y_final)
    outside = np.where(np.abs(y - y_final) > band)[0]
    if len(outside) == 0:
        return 0.0
    last_outside = outside[-1]
    if last_outside + 1 < len(t):
        return t[last_outside + 1]
    return t[-1]  # ยังไม่เข้า settling ภายในเวลาที่จำลอง


# ===================================================================
# การทดลองที่ 1: เปลี่ยนค่า K จาก 1 ถึง 10 (ครบทุกค่าจำนวนเต็ม, tau คงที่ = 1.0)
# ===================================================================
K_values = list(range(1, 11))   #  1,2,3,4,5,6,7,8,9,10 ครบทุกค่า
tau_fixed = 1.0

fig1, axes1 = plt.subplots(1, 2, figsize=(14, 5.5))
cmap_k = plt.cm.viridis(np.linspace(0, 1, len(K_values)))

# --- Open-loop plant response ---
for K, c in zip(K_values, cmap_k):
    y_out = simulate_open_loop(K, tau_fixed, u=1.0, t=t)
    axes1[0].plot(t, y_out, label=f'K = {K}', linewidth=2, color=c)

axes1[0].set_title('Open-Loop Plant Response\n(varying K, tau=1.0 )', fontsize=12, fontweight='bold')
axes1[0].set_xlabel('Time (minutes)')
axes1[0].set_ylabel('Temperature Change (°C)')
axes1[0].legend(loc='lower right', fontsize=8, ncol=2)
axes1[0].grid(True, alpha=0.3)
axes1[0].axhline(y=1, color='gray', linestyle='--', alpha=0.4)

# --- Closed-loop response ---
for K, c in zip(K_values, cmap_k):
    y_out = simulate_closed_loop(K, tau_fixed, Kp, setpoint, t=t)
    axes1[1].plot(t, y_out, label=f'K = {K}', linewidth=2, color=c)

axes1[1].set_title('Closed-Loop Response\n(varying K, tau=1.0)', fontsize=12, fontweight='bold')
axes1[1].set_xlabel('Time (minutes)')
axes1[1].set_ylabel('Temperature Change (°C)')
axes1[1].legend(loc='lower right', fontsize=8, ncol=2)
axes1[1].grid(True, alpha=0.3)
axes1[1].axhline(y=1, color='gray', linestyle='--', alpha=0.4)

plt.tight_layout()
plt.savefig(f'{OUT_DIR}/experiment1_vary_K.png', dpi=150)
plt.close()

print("=" * 70)
print("EXPERIMENT 1: Varying K from 1 to 10 (tau = 1.0 )")
print("=" * 70)
print(f"{'K':<6}{'Closed-loop Steady-State':<28}{'Settling Time (2%)':<22}{'Rise Time':<15}")
print("-" * 70)

exp1_rows = []
for K in K_values:
    y_out = simulate_closed_loop(K, tau_fixed, Kp, setpoint, t=t)
    y_final = y_out[-1]
    rt = compute_rise_time(t, y_out, y_final)
    st = compute_settling_time(t, y_out, y_final)
    print(f"{K:<6}{y_final:<28.4f}{st:<22.3f}{rt:<15.3f}")
    exp1_rows.append({
        'K': K,
        'tau': tau_fixed,
        'Kp': Kp,
        'Closed-loop Steady-State (°C)': y_final,
        'Settling Time 2% (min)': st,
        'Rise Time 10-90% (min)': rt,
    })

# ===================================================================
# การทดลองที่ 2: เพิ่ม tau เป็น 2 เท่า และ 3 เท่า (K คงที่ = 2.0)
# ===================================================================
K_fixed = 2.0
tau_values = [1.0, 2.0, 3.0]  # tau ฐาน, 2 เท่า, 3 เท่า
tau_labels = ['tau (base) = 1.0', 'tau x2 = 2.0', 'tau x3 = 3.0']

fig2, axes2 = plt.subplots(1, 2, figsize=(14, 5.5))

# --- Open-loop plant ---
for tau_v, label in zip(tau_values, tau_labels):
    y_out = simulate_open_loop(K_fixed, tau_v, u=1.0, t=t)
    axes2[0].plot(t, y_out, label=label, linewidth=2)

axes2[0].set_title('Open-Loop Plant Response\n(varying tau, K=2.0 )', fontsize=12, fontweight='bold')
axes2[0].set_xlabel('Time (minutes)')
axes2[0].set_ylabel('Temperature Change (°C)')
axes2[0].legend(loc='lower right', fontsize=9)
axes2[0].grid(True, alpha=0.3)

# --- Closed-loop ---
for tau_v, label in zip(tau_values, tau_labels):
    y_out = simulate_closed_loop(K_fixed, tau_v, Kp, setpoint, t=t)
    axes2[1].plot(t, y_out, label=label, linewidth=2)

axes2[1].set_title('Closed-Loop Response\n(varying tau, K=2.0)', fontsize=12, fontweight='bold')
axes2[1].set_xlabel('Time (minutes)')
axes2[1].set_ylabel('Temperature Change (°C)')
axes2[1].legend(loc='lower right', fontsize=9)
axes2[1].grid(True, alpha=0.3)

plt.tight_layout()
plt.savefig(f'{OUT_DIR}/experiment2_vary_tau.png', dpi=150)
plt.close()

print()
print("=" * 70)
print("EXPERIMENT 2: Varying tau (1x, 2x, 3x) with K = 2.0 fixed")
print("=" * 70)
print(f"{'tau':<8}{'Closed-loop Steady-State':<28}{'Settling Time (2%)':<22}{'Rise Time':<15}")
print("-" * 70)

exp2_rows = []
for tau_v, label in zip(tau_values, tau_labels):
    y_out = simulate_closed_loop(K_fixed, tau_v, Kp, setpoint, t=t)
    y_final = y_out[-1]
    rt = compute_rise_time(t, y_out, y_final)
    st = compute_settling_time(t, y_out, y_final)
    print(f"{tau_v:<8}{y_final:<28.4f}{st:<22.3f}{rt:<15.3f}")
    exp2_rows.append({
        'tau': tau_v,
        'tau label': label,
        'K': K_fixed,
        'Kp': Kp,
        'Closed-loop Steady-State (°C)': y_final,
        'Settling Time 2% (min)': st,
        'Rise Time 10-90% (min)': rt,
    })

# ===================================================================
# กราฟสรุปรวม: Combined comparison overview (K sweep)
# ===================================================================
fig3, ax3 = plt.subplots(figsize=(9, 6))
for K, c in zip(K_values, cmap_k):
    y_out = simulate_closed_loop(K, tau_fixed, Kp, setpoint, t=t)
    ax3.plot(t, y_out, label=f'K={K}', color=c, linewidth=2)
ax3.set_title('Summary: Closed-Loop Step Response for K = 1 to 10\n(tau = 1.0)', fontsize=13, fontweight='bold')
ax3.set_xlabel('Time (minutes)')
ax3.set_ylabel('Temperature Change (°C)')
ax3.legend(loc='lower right', fontsize=9, ncol=2)
ax3.grid(True, alpha=0.3)
plt.tight_layout()
plt.savefig(f'{OUT_DIR}/summary_K_sweep.png', dpi=150)
plt.close()

print()
print("All plots saved successfully. (ไม่มีการใช้ library 'control' เลย - ใช้ scipy.integrate.odeint แก้ ODE โดยตรง)")

# ===================================================================
# บันทึกผลลัพธ์เป็นไฟล์ .xlsx
# ===================================================================
df_exp1 = pd.DataFrame(exp1_rows)
df_exp2 = pd.DataFrame(exp2_rows)

# ---- raw time-series data (ทุกจุดเวลาจะใหญ่เกินไป -> สุ่มตัวอย่างทุกๆ 20 จุด ให้พอดูกราฟใน Excel ได้) ----
sample_idx = np.arange(0, len(t), 20)
t_sampled = t[sample_idx]

ts_data = {'Time (min)': t_sampled}
for K in K_values:
    y_out = simulate_closed_loop(K, tau_fixed, Kp, setpoint, t=t)
    ts_data[f'K={K}'] = y_out[sample_idx]
df_timeseries_K = pd.DataFrame(ts_data)

ts_data2 = {'Time (min)': t_sampled}
for tau_v, label in zip(tau_values, tau_labels):
    y_out = simulate_closed_loop(K_fixed, tau_v, Kp, setpoint, t=t)
    ts_data2[label] = y_out[sample_idx]
df_timeseries_tau = pd.DataFrame(ts_data2)

xlsx_path = f'{OUT_DIR}/control_simulation_results.xlsx'

with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
    df_exp1.to_excel(writer, sheet_name='Exp1_Vary_K', index=False)
    df_exp2.to_excel(writer, sheet_name='Exp2_Vary_tau', index=False)
    df_timeseries_K.to_excel(writer, sheet_name='TimeSeries_K', index=False)
    df_timeseries_tau.to_excel(writer, sheet_name='TimeSeries_tau', index=False)

# ---- จัดฟอร์แมตตารางให้ดูเป็นมืออาชีพ ----
wb = Workbook()
wb.remove(wb.active)
from openpyxl import load_workbook
wb = load_workbook(xlsx_path)

header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
body_font = Font(name='Arial', size=10)
thin = Side(style='thin', color='B0B0B0')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center')

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    max_col = ws.max_column
    max_row = ws.max_row

    # header row
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # body rows
    for row in range(2, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = body_font
            cell.alignment = center
            cell.border = border
            if isinstance(cell.value, float):
                cell.number_format = '0.0000'

    # column width auto-fit (approx)
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        max_len = max(
            [len(str(ws.cell(row=r, column=col).value)) for r in range(1, max_row + 1)] + [10]
        )
        ws.column_dimensions[letter].width = max_len + 3

    ws.freeze_panes = 'A2'

# ---- เพิ่มกราฟ Line Chart ในชีต TimeSeries ----
def add_line_chart(ws, n_series, n_rows, title, y_title):
    chart = LineChart()
    chart.title = title
    chart.style = 12
    chart.x_axis.title = 'Time (min)'
    chart.y_axis.title = y_title
    chart.width = 24
    chart.height = 12
    cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
    for col in range(2, n_series + 2):
        data = Reference(ws, min_col=col, min_row=1, max_row=n_rows + 1)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, get_column_letter(n_series + 3) + '2')

ws_k = wb['TimeSeries_K']
add_line_chart(ws_k, len(K_values), len(df_timeseries_K), 'Closed-Loop Response: Varying K', 'Temperature Change (°C)')

ws_tau = wb['TimeSeries_tau']
add_line_chart(ws_tau, len(tau_values), len(df_timeseries_tau), 'Closed-Loop Response: Varying tau', 'Temperature Change (°C)')

wb.save(xlsx_path)
print(f"\nบันทึกไฟล์ผลลัพธ์เป็นตาราง Excel เรียบร้อยแล้วที่: {xlsx_path}")
